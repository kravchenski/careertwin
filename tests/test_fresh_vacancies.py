"""The newest vacancies from every site: dates, Workday through its own JSON, the order, the workbook, the next steps."""

from __future__ import annotations

import io
import json
from datetime import date

import httpx
import pytest
from openpyxl import load_workbook

from careertwin import db
from careertwin.automation import freshness, guard
from careertwin.automation.excel import workbook
from careertwin.automation.models import Vacancy
from careertwin.automation.store import Store
from careertwin.jobs import workday
from careertwin.jobs.boards.base import to_iso_date

TODAY = date(2026, 10, 26)
SITE = "https://acme.wd3.myworkdayjobs.com/en-US/Careers/job/Gdansk/Python-Intern_R1"


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(guard, "DATA_DIR", tmp_path / "guard")
    value = Store()
    value.initialize()
    return value


def test_a_posting_date_is_read_from_whatever_the_source_gives():
    assert to_iso_date("2026-10-24T10:00:00Z") == "2026-10-24"
    assert to_iso_date(1_792_000_000_000, TODAY) == to_iso_date(1_792_000_000, TODAY) == "2026-10-14"
    assert to_iso_date("Posted Today", TODAY) == "2026-10-26" and to_iso_date("Posted Yesterday", TODAY) == "2026-10-25"
    assert to_iso_date("Posted 3 Days Ago", TODAY) == "2026-10-23" and to_iso_date("Posted 2 Weeks Ago", TODAY) == "2026-10-12"
    assert to_iso_date("", TODAY) == to_iso_date(None, TODAY) == to_iso_date("soon", TODAY) == ""


def test_the_site_is_read_from_any_workday_link():
    site = workday.parse_site(SITE)
    assert (site.host, site.tenant, site.name) == ("acme.wd3.myworkdayjobs.com", "acme", "Careers")
    assert site.api == "https://acme.wd3.myworkdayjobs.com/wday/cxs/acme/Careers"
    assert workday.parse_site("https://acme.wd3.myworkdayjobs.com/Careers").name == "Careers"  # no language in the address
    assert workday.parse_site("https://acme.wd3.myworkdayjobs.com/wday/cxs/acme/Careers/jobs").name == "Careers"
    assert workday.parse_site("https://example.com/Careers") is None
    assert workday.parse_site("https://evil.com/.wd3.myworkdayjobs.com/x") is None
    assert workday.parse_site("https://acme.wd3.myworkdayjobs.com/") is None


FACETS = [
    {"facetParameter": "timeType", "descriptor": "Time Type", "values": [{"descriptor": "Full time", "id": "ft", "count": 3}]},
    {
        "facetParameter": "locationMainGroup",
        "values": [
            {"facetParameter": "locationCountry", "descriptor": "Country", "values": [{"descriptor": "Germany", "id": "de", "count": 1}, {"descriptor": "Poland", "id": "pl", "count": 2}]},
            {"facetParameter": "locations", "descriptor": "City", "values": [{"descriptor": "Gdańsk, Poland", "id": "gdn", "count": 1}]},
        ],
    },
]


def test_a_country_or_a_city_is_found_in_the_filter_tree():
    assert workday.find_facet(FACETS, "Poland") == ("locationCountry", "pl")
    assert workday.find_facet(FACETS, "gdańsk") == ("locations", "gdn")
    assert workday.find_facet(FACETS, "Poland", parameter="locations") is None
    assert workday.find_facet(FACETS, "Narnia") is None and workday.find_facet([], "Poland") is None


def fake_workday(seen):
    postings = {
        "de": [{"title": "Java Intern", "externalPath": "/job/Berlin/Java-Intern_R3", "locationsText": "Berlin, Germany", "postedOn": "Posted Today"}],
        "pl": [
            {"title": "Python Intern", "externalPath": "/job/Gdansk/Python-Intern_R1", "locationsText": "Gdańsk, Poland", "postedOn": "Posted 3 Days Ago"},
            {"title": "Data Engineer Intern", "externalPath": "/job/Warsaw/Data-Intern_R2", "locationsText": "Warsaw, Poland", "postedOn": "Posted 30+ Days Ago"},
        ],
    }
    details = {
        "/job/Gdansk/Python-Intern_R1": {"title": "Python Intern", "startDate": "2026-10-22", "location": "Gdańsk, Poland", "remoteType": "Hybrid", "jobReqId": "R1", "externalUrl": SITE, "jobDescription": "<p>Python, <b>SQL</b></p>"},
        "/job/Warsaw/Data-Intern_R2": {"title": "Data Engineer Intern", "startDate": "2026-10-24", "location": "Warsaw, Poland", "jobReqId": "R2", "jobDescription": "<p>Spark</p>"},
    }

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append((request.method, request.url.path))
        if request.method == "POST":
            body = json.loads(request.content)
            facet = (body["appliedFacets"].get("locationCountry") or [""])[0]
            rows = postings.get(facet) or [row for group in postings.values() for row in group]
            return httpx.Response(200, json={"total": len(rows), "facets": FACETS, "jobPostings": rows})
        path = request.url.path.removeprefix("/wday/cxs/acme/Careers")
        return httpx.Response(200, json={"jobPostingInfo": details[path]})

    return httpx.AsyncClient(transport=httpx.MockTransport(handler))


async def test_the_newest_vacancies_of_a_workday_site_are_those_of_the_country_with_the_exact_day(monkeypatch):
    monkeypatch.setattr(workday, "PAUSE", 0)
    seen: list = []
    site = workday.parse_site(SITE)
    async with fake_workday(seen) as client:
        found = await workday.newest(site, "intern", place="Poland", limit=10, client=client, today=TODAY)
    assert [item["title"] for item in found] == ["Data Engineer Intern", "Python Intern"], "newest first, the German one left out"
    assert [item["posted_at"] for item in found] == ["2026-10-24", "2026-10-22"], "the start date, not '30+ days ago'"
    python = found[1]
    assert python["url"] == SITE and python["company"] == "Acme" and python["source"] == "workday"
    assert python["location"] == "Gdańsk, Poland (Hybrid)" and python["description"] == "Python, SQL" and python["external_id"] == "R1"
    async with fake_workday(seen) as client:
        recent = await workday.newest(site, "", place="Poland", limit=10, days=3, client=client, today=TODAY)
    assert [item["title"] for item in recent] == ["Data Engineer Intern"]
    async with fake_workday(seen) as client:
        assert await workday.newest(site, "", place="Narnia", limit=10, client=client, today=TODAY) == []


def add(store, url, title, *, posted="", company="Acme", location="Gdańsk", source="agent", score=None, state=None):
    job = store.add(Vacancy(url=url, title=title, company=company, location=location, source=source, posted_at=posted, description="d" * 600))
    if score is not None:
        store.update(job["id"], evaluation={"score": score, "category": "priority"}, state=state or "priority")
    return job["id"]


def test_the_sites_are_named_for_a_person():
    cases = {
        "https://www.linkedin.com/jobs/view/123": "LinkedIn", "https://pl.glassdoor.com/job-listing/x": "Glassdoor",
        "https://acme.wd3.myworkdayjobs.com/en-US/Careers/job/x": "Workday", "https://www.pracuj.pl/praca/x": "Pracuj.pl",
        "https://justjoin.it/job-offer/x": "Just Join IT", "https://boards.greenhouse.io/acme/jobs/1": "Greenhouse",
        "https://jobs.example.org/1": "jobs.example.org",
    }
    for url, label in cases.items():
        assert freshness.source_label(url) == label, url
    assert freshness.source_label("", "remoteok") == "Remoteok"


def test_the_newest_come_first_and_a_posting_without_a_date_is_dated_by_the_day_it_was_found(store):
    old = add(store, "https://x.example/1", "Old", posted="2026-10-01")
    mid = add(store, "https://www.linkedin.com/jobs/view/2", "Mid", posted="2026-10-23", score=71)
    new_ranked = add(store, "https://acme.wd3.myworkdayjobs.com/Careers/job/3", "New and ranked", posted="2026-10-26", score=88)
    new_plain = add(store, "https://x.example/4", "New, not ranked", posted="2026-10-26")
    found = add(store, "https://x.example/5", "Found today, no date")
    rows = freshness.newest(store, today=date.fromisoformat(store.job(found)["created_at"][:10]))
    order = [row["id"] for row in rows]
    assert order.index(new_ranked) < order.index(new_plain) < order.index(mid) < order.index(old), "the same day: the better fit first"
    undated = next(row for row in rows if row["id"] == found)
    assert undated["dated"] is False and undated["day"] == store.job(found)["created_at"][:10]
    linkedin = next(row for row in rows if row["id"] == mid)
    assert linkedin["source"] == "LinkedIn" and linkedin["dated"] is True and linkedin["score"] == 71
    recent = freshness.newest(store, days=3, today=date(2026, 10, 26))
    ids = {row["id"] for row in recent}
    assert {mid, new_ranked, new_plain} <= ids and old not in ids, "dated 3 days or less ago"
    assert len(freshness.newest(store, limit=2)) == 2


def test_the_workbook_opens_on_the_newest_sheet_with_a_link_and_a_posted_column(store):
    add(store, "https://www.linkedin.com/jobs/view/2", "Backend Intern", posted="2026-10-23", score=71)
    add(store, "https://x.example/9", "Junior Python", posted="2026-10-25")
    book = load_workbook(io.BytesIO(workbook(store)))
    assert book.sheetnames[:2] == ["Dashboard", "Applications"] and book.active.title == "Newest"
    sheet = book["Newest"]
    assert [cell.value for cell in sheet[1]][:6] == ["Posted", "Days ago", "Dated by the employer", "Site", "Company", "Role"]
    assert [sheet.cell(row=row, column=6).value for row in (2, 3)] == ["Junior Python", "Backend Intern"]
    assert sheet.cell(row=3, column=4).value == "LinkedIn" and sheet.cell(row=3, column=8).value == 71
    assert sheet.cell(row=3, column=11).hyperlink.target == "https://www.linkedin.com/jobs/view/2"
    skipped = add(store, "https://x.example/10", "Skipped one", posted="2026-10-26")
    store.update(skipped, state="skip")
    again = load_workbook(io.BytesIO(workbook(store)))
    assert [again["Newest"].cell(row=row, column=6).value for row in (2, 3, 4)] == ["Junior Python", "Backend Intern", None]
    assert "Skipped one" in [cell.value for cell in again["Applications"]["C"]], "still in the full log"
    assert book["Applications"].cell(row=1, column=14).value == "Posted"
    assert {book["Applications"].cell(row=row, column=14).value for row in (2, 3)} == {"2026-10-23", "2026-10-25"}


async def test_skipping_leaves_the_lists_and_an_application_that_went_out_cannot_be_skipped(store):
    from fastapi import HTTPException

    from careertwin.api.routes import automation

    job = add(store, "https://x.example/1", "Skip me", posted="2026-10-25")
    automation.engine.store = store
    skipped = await automation.skip(job)
    assert skipped["state"] == "skip" and store.events(job)[-1]["stage"] == "user_skip"
    sent = add(store, "https://x.example/2", "Sent", posted="2026-10-25")
    store.update(sent, state="submitted")
    with pytest.raises(HTTPException) as refused:
        await automation.skip(sent)
    assert refused.value.status_code == 409


async def test_a_workday_site_is_filed_through_the_api_and_a_bad_link_is_refused(store, monkeypatch):
    from fastapi import HTTPException

    from careertwin.api.routes import automation

    automation.engine.store = store

    async def fake(site, query, *, place, limit, days):
        assert (site.tenant, query, place, limit, days) == ("acme", "python", "Poland", 5, 14)
        return [{"url": SITE, "title": "Python Intern", "company": "Acme", "location": "Gdańsk, Poland", "description": "d" * 600,
                 "posted_at": "2026-10-22", "source": "workday", "external_id": "R1"}]

    monkeypatch.setattr(automation.workday, "newest", fake)
    result = await automation.workday_discover(automation.WorkdayRequest(url=SITE, query="python", limit=5, days=14))
    assert result["found"] == 1 and result["site"] == "acme/Careers"
    job = store.job(result["jobs"][0])
    assert job["vacancy"]["posted_at"] == "2026-10-22" and job["vacancy"]["source"] == "workday" and job["state"] == "new"
    with pytest.raises(HTTPException) as refused:
        await automation.workday_discover(automation.WorkdayRequest(url="https://example.com/jobs"))
    assert refused.value.status_code == 422


async def test_the_agent_tools_pass_the_date_and_ask_for_the_newest(monkeypatch):
    from careertwin.agentic import tools

    calls = []

    class Client:
        async def post(self, path, body=None, **options):
            calls.append(("post", path, body))
            return {"id": "j1", "vacancy": {"title": body.get("title")}, "state": "new"}

        async def get(self, path, **params):
            calls.append(("get", path, params))
            return []

    monkeypatch.setattr(tools, "_client", lambda: Client())
    await tools.import_vacancy("https://x.example/1", "Intern", posted_at="2026-10-24")
    await tools.discover_workday(SITE, "python", limit=5)
    await tools.newest_vacancies(days=7, limit=10)
    assert calls[0][2]["posted_at"] == "2026-10-24"
    assert calls[1][:2] == ("post", "/v1/automation/workday") and calls[1][2]["place"] == "Poland"
    assert calls[2] == ("get", "/v1/automation/newest", {"days": 7, "limit": 10})
