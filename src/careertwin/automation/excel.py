from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
from pathlib import Path
from typing import Any

from careertwin.db import db_conn
from careertwin.profile.context import context_path, read_context
from careertwin.settings import OUTPUT_DIR
from careertwin.vault import redact, redact_data

from . import freshness
from .mail import MailStore
from .monitor import WatchStore
from .stats import stats
from .store import Store


class ExcelAgent:
    """Keep a materialized workbook alongside the durable SQLite event journal."""

    def __init__(self, store: Store):
        self.store = store
        self.path = OUTPUT_DIR / "CareerTwin-applications.xlsx"
        self.error = ""
        self.revision = None

    async def flush(self) -> None:
        with db_conn() as db:
            revision = (
                db.execute("SELECT MAX(id) FROM automation_events").fetchone()[0],
                db.execute("SELECT MAX(updated_at) FROM automation_jobs").fetchone()[0],
                hashlib.sha256(read_context().encode()).hexdigest(),
            )
        if revision == self.revision and self.path.is_file():
            return

        def write():
            self.path.parent.mkdir(parents=True, exist_ok=True)
            temporary = Path(str(self.path) + ".tmp")
            temporary.write_bytes(workbook(self.store))
            temporary.chmod(0o600)  # it holds the answers, the history and the whole ABOUT.md
            os.replace(temporary, self.path)

        try:
            writing = asyncio.create_task(asyncio.to_thread(write))
            try:
                await asyncio.shield(writing)
            except asyncio.CancelledError:
                await writing
                raise
            self.revision, self.error = revision, ""
        except Exception as exc:
            self.error = str(exc)

    async def run(self) -> None:
        while True:
            await self.flush()
            await asyncio.sleep(2)


def safe(value: Any) -> str:
    """A cell that a spreadsheet would run as a formula is written as text (OWASP: = + - @, tab, CR)."""
    value = redact(str(value if value is not None else ""))
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def chunks(value: Any) -> list[str]:
    text = json.dumps(redact_data(value), ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)
    return [safe(text[i : i + 30000]) for i in range(0, len(text), 30000)] or [""]


def workbook(store: Store) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    dashboard = book.active
    dashboard.title = "Dashboard"
    fill_dashboard(dashboard, stats(store))
    jobs = book.create_sheet("Applications")
    jobs.append(
        [
            "ID",
            "Company",
            "Role",
            "URL",
            "Source",
            "Fit score",
            "Confidence",
            "Category",
            "Status",
            "CV path",
            "Created",
            "Updated",
            "Error",
            "Posted",
        ]
    )
    for job in store.jobs():
        vacancy, evaluation, package = job["vacancy"], job["evaluation"] or {}, job["package"] or {}
        jobs.append(
            [
                safe(value)
                for value in [
                    job["id"],
                    vacancy["company"],
                    vacancy["title"],
                    vacancy["url"],
                    vacancy["source"],
                    evaluation.get("score", ""),
                    evaluation.get("confidence", ""),
                    evaluation.get("category", ""),
                    job["state"],
                    package.get("cv_path", ""),
                    job["created_at"],
                    job["updated_at"],
                    job["error"],
                    vacancy.get("posted_at", ""),
                ]
            ]
        )
    fresh = book.create_sheet("Newest")
    fresh.append(["Posted", "Days ago", "Dated by the employer", "Site", "Company", "Role", "Location", "Fit", "Category", "Status", "Link"])
    for row in freshness.newest(store, limit=2000):
        if row["state"] in freshness.HIDDEN:  # skipped or closed: they stay in the Applications sheet
            continue
        fresh.append(
            [
                row["day"],
                row["age_days"] if row["age_days"] is not None else "",
                "yes" if row["dated"] else "found on this day",
                *[safe(row[key]) for key in ("source", "company", "title", "location")],
                row["score"] if row["score"] is not None else "",
                row["category"],
                row["state"],
                safe(row["url"]),
            ]
        )
        link = fresh.cell(row=fresh.max_row, column=11)
        if str(row["url"]).startswith(("http://", "https://")):
            link.hyperlink = row["url"]
    book.active = book.sheetnames.index("Newest")  # the file opens on the list to look at first
    answers = book.create_sheet("Answers")
    answers.append(["Job ID", "Date", "Field", "Value part", "Exact value", "Evidence"])
    history = book.create_sheet("History")
    history.append(["Event ID", "Job ID", "Date", "Stage", "Part", "Full event JSON"])
    for event in store.events():
        payload = event["payload"]
        if event["stage"] == "browser_action" and payload.get("field"):
            values = chunks(payload.get("value", ""))
            evidence = chunks(payload.get("evidence", ""))
            for index in range(max(len(values), len(evidence))):
                answers.append(
                    [
                        event["job_id"],
                        event["created_at"],
                        safe(payload["field"]),
                        index + 1,
                        values[index] if index < len(values) else "",
                        evidence[index] if index < len(evidence) else "",
                    ]
                )
        for index, part in enumerate(chunks(payload), 1):
            history.append(
                [event["id"], event["job_id"], event["created_at"], event["stage"], index, part]
            )
    snapshots = book.create_sheet("Snapshots")
    snapshots.append(["Job ID", "Part", "Full vacancy, evaluation and document package JSON"])
    for job in store.jobs():
        for index, part in enumerate(
            chunks({key: job[key] for key in ("vacancy", "evaluation", "package")}), 1
        ):
            snapshots.append([job["id"], index, part])
    files = book.create_sheet("Files")
    files.append(["Job ID", "Kind", "Absolute path", "SHA256", "Date"])
    for item in store.files():
        files.append(
            [safe(item[key]) for key in ("job_id", "kind", "path", "sha256", "created_at")]
        )
    known = book.create_sheet("Saved answers")
    known.append(["Job ID", "Question", "Part", "User answer"])
    for job_id in ["", *[job["id"] for job in store.jobs()]]:
        for question, answer in store.answers(job_id).items():
            for index, part in enumerate(chunks(answer), 1):
                known.append([job_id, safe(question), index, part])
    letters = book.create_sheet("Mail")
    letters.append(["Date", "Direction", "Kind", "State", "Job ID", "From", "To", "Subject", "Summary"])
    for row in reversed(MailStore().list(limit=5000)):  # bodies stay out: the log is for counting, not for reading
        letters.append([safe(value) for value in (row["created_at"], row["direction"], row["kind"], row["state"], row["job_id"], row["from_addr"], row["to_addr"], row["subject"], row["summary"])])
    companies = book.create_sheet("Companies")
    companies.append(["Company", "Careers page", "ATS", "Open vacancies", "Last visit", "Last success", "Failures in a row", "Last error"])
    for row in WatchStore().targets(enabled_only=False):
        companies.append([safe(row["name"]), safe(row["careers_url"]), safe(row["ats"]), row["open_count"], safe(row["last_checked"] or ""), safe(row["last_ok"] or ""), row["failures"], safe(row["last_error"])])
    shared = book.create_sheet("Candidate context")
    shared.append(["ABOUT.md absolute path", "SHA256", "Part", "Current shared context"])
    content = redact(read_context())
    checksum = hashlib.sha256(content.encode()).hexdigest()
    for index, part in enumerate(chunks(content), 1):
        shared.append([safe(str(context_path().resolve())), checksum, index, part])
    for sheet in book:
        if sheet.title == "Dashboard":
            continue
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                65, max(15, len(str(column[0].value)) + 5)
            )
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def fill_dashboard(sheet, numbers: dict[str, Any]) -> None:
    """The first sheet: how the search is going, in numbers and one chart."""
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font

    bold = Font(bold=True)

    def put(label: str, value: Any = "") -> None:
        sheet.append([safe(label), value if isinstance(value, (int, float)) else safe(value)])
        sheet.cell(sheet.max_row, 1).font = bold if value == "" else Font()

    put("CareerTwin dashboard", "")
    sheet.cell(1, 1).font = Font(bold=True, size=14)
    sheet.append([])
    put("Vacancies in the pipeline", numbers["jobs"])
    put("Applied", numbers["applied"])
    put("Received a reply", numbers["replied"])
    put("Reply rate", "" if numbers["reply_rate"] is None else f"{numbers['reply_rate']:.0%}")
    put("Median days to the first reply", "" if numbers["median_days_to_reply"] is None else numbers["median_days_to_reply"])
    put("Interviews, offers, hired", numbers["interviews"])
    put("Rejections", numbers["rejections"])
    put("Submitted today", numbers["submitted_today"])
    put("Companies watched / failing", f"{numbers['monitor']['watching']} / {numbers['monitor']['failing']}")
    put("Mail received / sent / drafts waiting", f"{numbers['mail']['received']} / {numbers['mail']['sent']} / {numbers['mail']['drafts']}")
    sheet.append([])
    put("By state")
    for state, count in numbers["by_state"].items():
        put(state, count)
    sheet.append([])
    put("By source")
    sheet.append(["Source", "Found", "Applied"])
    for source, row in numbers["by_source"].items():
        sheet.append([safe(source), row["found"], row["applied"]])
    sheet.append([])
    put("Last days")
    header = sheet.max_row + 1
    sheet.append(["Date", "Found", "Submitted", "Replies"])
    for day in numbers["timeline"]:
        sheet.append([day["date"], day["found"], day["submitted"], day["replies"]])
    chart = BarChart()
    chart.title, chart.height, chart.width = "Found, submitted and replies per day", 8, 18
    chart.add_data(Reference(sheet, min_col=2, max_col=4, min_row=header, max_row=sheet.max_row), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=header + 1, max_row=sheet.max_row))
    sheet.add_chart(chart, "E3")
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 18
