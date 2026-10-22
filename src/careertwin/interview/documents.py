"""Plain text from the files you attach to an interview: CVs as PDF, DOCX or text."""

from __future__ import annotations

import re
from pathlib import Path

MIN_TEXT = 120


class DocumentError(ValueError):
    """The file is missing, unsupported or has no readable text."""


def read_document(path: str | Path) -> str:
    file = Path(path).expanduser()
    if not file.is_file():
        raise DocumentError(f"{file} does not exist")
    suffix = file.suffix.lower()
    if suffix == ".pdf":
        text = _pdf_text(file)
    elif suffix == ".docx":
        text = _docx_text(file)
    elif suffix in {".txt", ".md", ""}:
        text = file.read_text(encoding="utf-8", errors="replace")
    else:
        raise DocumentError(f"{file.name}: unsupported format {suffix}; use PDF, DOCX or text")
    text = re.sub(r"[ \t]+\n", "\n", re.sub(r"\n{3,}", "\n\n", text)).strip()
    if len(text) < MIN_TEXT:
        raise DocumentError(
            f"{file.name} has almost no readable text ({len(text)} characters). "
            "A scanned or blank PDF cannot be used; attach the real CV."
        )
    return text


def _pdf_text(file: Path) -> str:
    candidates: list[str] = []
    try:
        from pdfminer.high_level import extract_text

        candidates.append(extract_text(str(file)) or "")
    except Exception:
        pass
    if not candidates or len(candidates[0].strip()) < MIN_TEXT:
        try:
            from pypdf import PdfReader

            candidates.append("\n".join(page.extract_text() or "" for page in PdfReader(str(file)).pages))
        except Exception:
            pass
    if not candidates:
        raise DocumentError(f"Could not read {file.name}")
    return max(candidates, key=lambda text: len(text.strip()))


def _docx_text(file: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentError("DOCX needs python-docx: uv sync") from exc
    document = Document(str(file))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    return "\n".join(parts)
