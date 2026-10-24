"""The Excel agent keeps a dashboard, a mail log and the health of the watched companies; the same numbers are in the API."""

from __future__ import annotations

import io
import os
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from careertwin import db
from careertwin.automation import guard
from careertwin.automation.excel import ExcelAgent, safe, workbook
from careertwin.automation.mail import MailStore
from careertwin.automation.models import Settings, Vacancy
from careertwin.automation.monitor import Target, WatchStore
from careertwin.automation.stats import attention, stats
from careertwin.automation.store import Store

PROFILE = "## Contact\n| Full name | Alex Demo |\n| Email | alex@example.com |\n## Experience\n- Built Python APIs.\n"
SECRET = "hunter2-hunter2-hunter2"


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(guard, "DATA_DIR", tmp_path / "guard")
    value = Store()
    value.initialize()
    value.save_settings(Settings(context=PROFILE))
    return value


def at(days_ago: float) -> str:
    return (datetime.now(timezone.utc) - timedelta(days=days_ago)).isoformat()


def add(store, title, state, source="remotive", company="Acme"):
    job = store.add(Vacancy(url=f"https://{company.lower()}.example/jobs/{title.replace(' ', '-')}", title=title, company=company, source=source))
    store.update(job["id"], state=state)
    return store.job(job["id"])


def stamp(stage, job_id, when):
    with db.db_conn() as conn:
        conn.execute("UPDATE automation_events SET created_at=? WHERE stage=? AND job_id=?", (when, stage, job_id))
        conn.commit()


@pytest.fixture
def busy(store):
    a = add(store, "Python Developer", "submitted", source="company:Acme")
    b = add(store, "Backend Engineer", "rejected", company="Globex")
    c = add(store, "Data Engineer", "new", company="Initech")
    d = add(store, "API Developer", "form_ready", company="Hooli")
    e = add(store, "SRE", "awaiting_answers", company="Umbrella")
    add(store, "Platform", "closed", company="Stark")
    store.event(a["id"], "submit_attempt", {})
    store.event(a["id"], "submission_confirmed", {})
    store.event(b["id"], "submission_confirmed", {})
    stamp("submission_confirmed", a["id"], at(5))
    stamp("submission_confirmed", b["id"], at(4))
    mails = MailStore()
    mails.add(job_id=a["id"], direction="in", kind="interview_invite", state="new", from_addr="hr@acme.example", subject="Interview?", summary="Asks for a call.", message_id="<1>", created_at=at(3))
    mails.add(job_id=b["id"], direction="in", kind="rejection", state="handled", from_addr="hr@globex.example", subject="Update", message_id="<2>", created_at=at(3))
    mails.add(job_id=c["id"], direction="out", kind="follow_up", state="draft", to_addr="jobs@initech.example", subject="Following up", body="Dear team, a follow-up.", message_id="")
    mails.add(job_id=a["id"], direction="out", kind="application", state="sent", to_addr="jobs@acme.example", subject="Application", message_id="<out1>")
    watch = WatchStore()
    watch.add(Target("Acme", "https://careers.acme.example"))
    watch.add(Target("Globex", ""))
    watch.health("Acme", ok=True, interval=timedelta(hours=1), open_count=7)
    watch.health("Globex", ok=False, interval=timedelta(hours=1), error="down")
    return {"a": a, "b": b, "c": c, "d": d, "e": e}


def test_the_numbers_of_a_search(store, busy):
    numbers = stats(store)
    assert numbers["jobs"] == 6 and numbers["by_state"]["new"] == 1 and numbers["by_state"]["closed"] == 1
    assert numbers["applied"] == 2 and numbers["interviews"] == 0 and numbers["rejections"] == 1
    assert numbers["by_source"]["company:Acme"] == {"found": 1, "applied": 1} and numbers["by_source"]["remotive"] == {"found": 5, "applied": 1}
    assert numbers["replied"] == 2 and numbers["reply_rate"] == 1.0
    assert numbers["median_days_to_reply"] == 1.5  # replies came 2 and 1 days after the applications
    assert numbers["submitted_today"] == 1
    assert numbers["mail"] == {"received": 2, "sent": 1, "drafts": 1, "to_read": 1}
    assert numbers["monitor"] == {"watching": 2, "failing": 1, "open": 7}
    assert len(numbers["timeline"]) == 14 and numbers["timeline"][-1]["submitted"] == 1 and numbers["timeline"][-1]["found"] == 6
    assert sum(day["replies"] for day in numbers["timeline"]) == 2


def test_an_empty_pipeline_has_no_rates_and_no_errors(store):
    numbers = stats(store)
    assert numbers["jobs"] == 0 and numbers["reply_rate"] is None and numbers["median_days_to_reply"] is None
    assert attention(store) == []


def test_what_waits_for_the_user_comes_most_urgent_first(store, busy):
    items = attention(store)
    kinds = [item["kind"] for item in items]
    assert kinds == ["interview_invite", "submit", "email", "answer"]
    by_kind = {item["kind"]: item for item in items}
    assert by_kind["submit"]["job_id"] == busy["d"]["id"] and by_kind["submit"]["actions"] == ["submit", "skip"]
    assert by_kind["email"]["actions"] == ["approve", "reject"] and "jobs@initech.example" in by_kind["email"]["title"]
    assert by_kind["interview_invite"]["detail"] == "Asks for a call." and "Python Developer at Acme" in by_kind["interview_invite"]["job"]
    assert by_kind["answer"]["actions"] == ["answer"]


def test_a_handled_message_and_a_sent_draft_leave_the_list(store, busy):
    mails = MailStore()
    for row in mails.list():
        if row["state"] in {"new", "draft"}:
            mails.update(row["id"], state="handled" if row["direction"] == "in" else "sent")
    assert [item["kind"] for item in attention(store)] == ["submit", "answer"]


def test_the_workbook_starts_with_the_dashboard_and_keeps_its_old_sheets(store, busy):
    book = load_workbook(io.BytesIO(workbook(store)))
    assert book.sheetnames[:2] == ["Dashboard", "Applications"]
    assert {"Mail", "Companies", "Answers", "History", "Files", "Candidate context"} <= set(book.sheetnames)
    dashboard = book["Dashboard"]
    cells = {row[0].value: row[1].value for row in dashboard.iter_rows(min_row=1, max_col=2) if row[0].value}
    assert cells["Applied"] == 2 and cells["Reply rate"] == "100%" and cells["Median days to the first reply"] == 1.5
    assert cells["Companies watched / failing"] == "2 / 1" and cells["Mail received / sent / drafts waiting"] == "2 / 1 / 1"
    assert len(dashboard._charts) == 1
    assert book["Applications"]["I2"].value in {"new", "submitted", "rejected", "form_ready", "awaiting_answers", "closed"}


def test_the_mail_and_company_sheets_carry_no_message_bodies(store, busy):
    book = load_workbook(io.BytesIO(workbook(store)))
    letters = [[cell.value for cell in row] for row in book["Mail"].iter_rows(min_row=2)]
    assert len(letters) == 4 and "Dear team" not in str(letters)
    companies = {row[0].value: [cell.value for cell in row] for row in book["Companies"].iter_rows(min_row=2)}
    assert companies["Acme"][3] == 7 and companies["Globex"][6] == 1 and companies["Globex"][7] == "down"


@pytest.mark.parametrize("value", ["=HYPERLINK(\"http://evil\",\"x\")", "+1+1", "-2+3", "@SUM(1)", "\t=1+1", "\r=1+1"])
def test_a_cell_that_would_run_as_a_formula_is_written_as_text(value):
    assert safe(value).startswith("'") and safe(value)[1:] == value


def test_a_secret_never_reaches_the_workbook(store, busy, monkeypatch):
    monkeypatch.setenv("ACCOUNT_PASSWORD", SECRET)
    store.save_settings(Settings(context=PROFILE + f"\n## Notes\n- my password is {SECRET}\n"))
    add(store, f"Dev {SECRET}", "new", company=f"Co {SECRET}")
    book = load_workbook(io.BytesIO(workbook(store)))
    everything = " ".join(str(cell.value) for sheet in book for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert SECRET not in everything and "[secret:ACCOUNT_PASSWORD]" in everything


async def test_the_workbook_file_is_private_to_its_owner(store, busy, tmp_path):
    tracker = ExcelAgent(store)
    tracker.path = tmp_path / "out" / "applications.xlsx"
    await tracker.flush()
    assert tracker.path.is_file() and not tracker.error
    if os.name == "posix":
        assert oct(tracker.path.stat().st_mode & 0o777) == "0o600"


@pytest.fixture
def api(store, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from careertwin.api.routes import automation
    from careertwin.automation.engine import Engine

    monkeypatch.setattr(automation, "engine", Engine(store))
    app = FastAPI()
    app.include_router(automation.router)
    with TestClient(app) as client:
        yield client


def test_the_api_serves_the_summary_and_the_attention_list(api, busy):
    summary = api.get("/v1/automation/summary").json()
    assert summary["applied"] == 2 and summary["monitor"]["failing"] == 1
    items = api.get("/v1/automation/attention").json()
    assert [item["kind"] for item in items] == ["interview_invite", "submit", "email", "answer"]


def test_the_agent_tools_read_both(busy):
    from careertwin.agentic.registry import get_tool

    for name in ("get_summary", "what_needs_me"):
        tool = get_tool(name)
        assert not tool.mutating and not tool.outward
