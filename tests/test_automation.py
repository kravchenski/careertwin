"""Contract and recovery checks for the four-agent workflow; no real applications."""

from __future__ import annotations

import asyncio
import io
import json
from pathlib import Path
from unittest.mock import AsyncMock

import httpx
import pytest

from careertwin import db
from careertwin.automation import browser, documents
from careertwin.automation.engine import Engine
from careertwin.automation.excel import ExcelAgent, workbook
from careertwin.automation.models import Settings, Vacancy
from careertwin.automation.ranker import DIMENSIONS, GATES, JevRanker
from careertwin.automation.sources import parse_posting
from careertwin.automation.store import Store

PROFILE = """## Contact
| Full name | Test Candidate |
| Email | test@example.com |
| Location | Gdańsk, Poland |
## Experience
### Developer, Freelance | 2024 to present
- Built Python APIs with FastAPI.
- Deployed services with Docker and Linux.
- Wrote PostgreSQL queries and indexes.
## Languages
| English | B2+ (daily technical communication). Never write C1. |
## Education
- BSc Computer Science, in progress, expected June 2028.
## Applications log
- Already submitted a separate application.
"""


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    value = Store()
    value.initialize()
    value.save_settings(Settings(context=PROFILE))
    return value


@pytest.fixture
def vacancy():
    return Vacancy(
        url="https://example.com/jobs/python",
        title="Python Developer",
        company="Example",
        description=(
            "Build APIs with Python and FastAPI, deploy with Docker, and "
            "use PostgreSQL. Work remotely from Poland. "
        )
        * 6,
    )


def jev_response(**gate_overrides):
    answers = {name: {"score": 4, "confidence": 0.95} for name in DIMENSIONS}
    answers.update({name: {"choice": "satisfied", "confidence": 0.95} for name in GATES})
    for name, value in gate_overrides.items():
        answers[name] = value
    return {"model": "jev-test", "answers": answers}


def judge(payload, monkeypatch):
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-only")

    def handle(request):
        assert request.url == "https://api.typesafe.ai/v1/systemone"
        assert request.headers["authorization"] == "Bearer test-only"
        body = json.loads(request.content)
        assert body["questions"]
        return httpx.Response(200, json=payload)

    return JevRanker(httpx.MockTransport(handle))


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "overrides,category",
    [
        ({}, "priority"),
        ({"authorization": {"choice": "gap", "confidence": 0.95}}, "skip"),
        ({"location": {"choice": "unknown", "confidence": 0.95}}, "needs_info"),
        ({"language": {"choice": "satisfied", "confidence": 0.4}}, "needs_info"),
        ({"education": {"choice": "gap", "confidence": 0.95}}, "stretch"),
    ],
)
async def test_jev_eligibility_gates(vacancy, monkeypatch, overrides, category):
    result = await judge(jev_response(**overrides), monkeypatch).evaluate(
        vacancy, Settings(context=PROFILE, decision_backend="typesafe")
    )
    assert result["score"] == 100
    assert result["category"] == category
    assert "not probability" in result["meaning"]


@pytest.mark.asyncio
async def test_no_fake_ranking_without_key(vacancy, monkeypatch):
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    with pytest.raises(RuntimeError, match="TYPESAFE_API_KEY"):
        await JevRanker().evaluate(vacancy, Settings(decision_backend="typesafe"))


@pytest.mark.asyncio
async def test_invalid_jev_response_is_rejected(vacancy, monkeypatch):
    payload = jev_response()
    payload["answers"]["skills"]["score"] = 100
    with pytest.raises(ValueError, match="Invalid Jev score"):
        await judge(payload, monkeypatch).evaluate(vacancy, Settings(decision_backend="typesafe"))


@pytest.mark.asyncio
async def test_semantic_answer_verification(monkeypatch):
    evaluator = judge({"answers": {"supported": {"choice": "no", "confidence": 0.99}}}, monkeypatch)
    assert not await evaluator.supports_answer(
        "English level?", "C1", "English B2+", PROFILE, {}, Settings(decision_backend="typesafe")
    )


def test_linkedin_tracking_deduplication(store):
    first = store.add(
        Vacancy(
            url=(
                "https://www.linkedin.com/jobs/search-results/?currentJobId=4"
                "422038307&keywords=python&refId=abc"
            ),
            title="Python",
        )
    )
    second = store.add(
        Vacancy(url="https://www.linkedin.com/jobs/view/4422038307/?utm_source=foo", title="Python")
    )
    assert first["id"] == second["id"]
    assert len(store.jobs()) == 1


@pytest.mark.parametrize(
    "interrupted,expected",
    [
        ("submitting", "submission_uncertain"),
        ("verifying", "submission_uncertain"),
        ("filling", "needs_review"),
    ],
)
def test_restart_recovery(store, vacancy, interrupted, expected):
    job = store.add(vacancy)
    store.update(job["id"], state=interrupted)
    store.initialize()
    assert store.job(job["id"])["state"] == expected
    assert store.events(job["id"])[-1]["payload"]["previous_state"] == interrupted


def test_user_answers_are_scoped_to_job(store, vacancy):
    first = store.add(vacancy)
    second = store.add(vacancy.model_copy(update={"url": "https://example.com/jobs/other"}))
    store.answer("Salary?", "10000 PLN", first["id"])
    assert store.answers(first["id"])["Salary?"] == "10000 PLN"
    assert store.answers(second["id"]) == {}


def test_external_ids_from_different_feeds_do_not_collide(store, vacancy):
    first = store.add(vacancy.model_copy(update={"source": "remotive", "external_id": "123"}))
    second = store.add(
        vacancy.model_copy(
            update={"source": "remoteok", "external_id": "123", "url": "https://example.com/other"}
        )
    )
    assert first["id"] != second["id"]


def test_full_jsonld_description_is_preserved(vacancy):
    description = "A long detailed requirement. " * 400
    html = (
        '<script type="application/ld+json">'
        + json.dumps(
            {
                "@graph": [
                    {
                        "@type": "JobPosting",
                        "title": "Engineer",
                        "description": f"<p>{description}</p>",
                        "hiringOrganization": {"name": "Real company"},
                    }
                ]
            }
        )
        + "</script>"
    )
    parsed = parse_posting(html, vacancy)
    assert parsed.company == "Real company"
    assert parsed.description == description.strip()


@pytest.mark.asyncio
async def test_pdf_uses_selected_facts_and_is_readable(store, vacancy, tmp_path, monkeypatch):
    from pypdf import PdfReader

    monkeypatch.setattr(documents, "OUTPUT_DIR", tmp_path / "outputs")
    monkeypatch.setattr(
        documents,
        "call_ai",
        AsyncMock(return_value='{"fact_ids":["fact_1","fact_2","fact_3","fact_4","fact_5"]}'),
    )
    job = store.add(vacancy)
    package = await documents.DocumentAgent().prepare(job, store.settings(), store)
    pdf = Path(package["cv_path"])
    assert pdf.read_bytes().startswith(b"%PDF")
    extracted = "\n".join(page.extract_text() for page in PdfReader(pdf).pages)
    assert "test@example.com" in extracted and "FastAPI" in extracted and "B2+" in extracted
    assert "C1" not in extracted and "Already submitted" not in extracted
    assert len(store.files(job["id"])) == 6
    assert package["checks"]["pages"] <= 2


@pytest.mark.asyncio
async def test_cv_cannot_invent_fact_ids(store, vacancy, monkeypatch):
    monkeypatch.setattr(
        documents, "call_ai", AsyncMock(return_value='{"fact_ids":["fact_1","fact_2","invented"]}')
    )
    with pytest.raises(ValueError, match="unknown facts"):
        await documents.DocumentAgent().prepare(store.add(vacancy), store.settings(), store)


class FakeMCP:
    page = '[ref_1] textbox "English level"\n[ref_2] button "Submit application"'
    clicked_page = "Thank you for applying"
    tab_url = "https://example.com/jobs/python"  # where the tab is; the vacancy's own site
    instances = []

    def __init__(self):
        self.calls = []
        self.definitions = {
            name: {"name": name, "inputSchema": {}}
            for name in [*browser.ALLOWED, "tab_create", "create_file", "tabs_list"]
        }
        self.__class__.instances.append(self)

    async def connect(self):
        pass

    async def close(self):
        pass

    async def tool(self, name, args):
        self.calls.append((name, args))
        if name == "tabs_list":
            value = json.dumps({"currentTabId": 123, "tabs": [{"id": 123, "title": "t", "url": self.tab_url}]})
        elif name == "tab_create":
            value = json.dumps({"tabId": 123, "page": {"text": self.page}})
        elif name == "click":
            value = self.clicked_page
        else:
            value = self.page
        return {"content": [{"type": "text", "text": value}]}


@pytest.fixture
def form(store, vacancy, monkeypatch):
    FakeMCP.instances = []
    monkeypatch.setattr(browser, "MCP", FakeMCP)
    job = store.add(vacancy)
    store.update(job["id"], state="prepared", package={"profile": PROFILE})
    return store.job(job["id"])


@pytest.mark.asyncio
async def test_fill_stops_before_submission(store, form, monkeypatch):
    monkeypatch.setattr(
        browser,
        "call_ai",
        AsyncMock(return_value='{"tool":"click","args":{"ref":"ref_2"},"submit":true}'),
    )
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert store.job(form["id"])["state"] == "form_ready"
    assert not any(name == "click" for name, _ in FakeMCP.instances[0].calls)
    assert store.submissions_today() == 0


@pytest.mark.asyncio
async def test_submission_requires_observed_confirmation(store, form, monkeypatch):
    monkeypatch.setattr(
        browser,
        "call_ai",
        AsyncMock(return_value='{"tool":"click","args":{"ref":"ref_2"},"submit":true}'),
    )
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=True)
    assert store.job(form["id"])["state"] == "submitted"
    assert store.submissions_today() == 1
    assert store.events(form["id"])[-1]["stage"] == "submission_confirmed"


@pytest.mark.asyncio
async def test_unknown_submission_is_verified_without_another_click(store, form, monkeypatch):
    store.update(form["id"], state="submission_uncertain", tab_id=123)
    model = AsyncMock()
    monkeypatch.setattr(browser, "call_ai", model)
    await browser.BrowserAgent().apply(
        store.job(form["id"]), store.settings(), store, allow_submit=True
    )
    assert store.job(form["id"])["state"] == "submission_uncertain"
    model.assert_not_called()
    assert not any(name == "click" for name, _ in FakeMCP.instances[0].calls)
    engine = Engine(store)
    with pytest.raises(ValueError, match="Verify"):
        engine.enqueue(form["id"], "rank")


@pytest.mark.asyncio
async def test_unsupported_answer_pauses_without_typing(store, form, monkeypatch):
    monkeypatch.setattr(
        browser,
        "call_ai",
        AsyncMock(
            return_value=json.dumps(
                {
                    "tool": "type",
                    "args": {"ref": "ref_1", "text": "C1"},
                    "field": "English level",
                    "value": "C1",
                    "evidence": "B2+",
                }
            )
        ),
    )
    monkeypatch.setattr(browser.JevRanker, "supports_answer", AsyncMock(return_value=False))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=True)
    assert store.job(form["id"])["state"] == "awaiting_answers"
    assert not any(name == "type" for name, _ in FakeMCP.instances[0].calls)
    assert store.events(form["id"])[-1]["payload"]["questions"] == ["English level"]


@pytest.mark.asyncio
async def test_missing_evidence_prevents_typing(store, form, monkeypatch):
    monkeypatch.setattr(
        browser,
        "call_ai",
        AsyncMock(
            return_value=(
                '{"tool":"type","args":{"ref":"ref_1","text":"C1"},"field":"E'
                'nglish","value":"C1","evidence":"invented C1"}'
            )
        ),
    )
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=True)
    assert store.job(form["id"])["state"] == "needs_review"
    assert not any(name == "type" for name, _ in FakeMCP.instances[0].calls)


@pytest.mark.asyncio
async def test_cancellation_after_submit_preserves_uncertainty(store, form, monkeypatch):
    class Interrupted(FakeMCP):
        async def tool(self, name, args):
            if name == "click":
                raise asyncio.CancelledError()
            return await super().tool(name, args)

    monkeypatch.setattr(browser, "MCP", Interrupted)
    monkeypatch.setattr(
        browser,
        "call_ai",
        AsyncMock(return_value='{"tool":"click","args":{"ref":"ref_2"},"submit":true}'),
    )
    engine = Engine(store)
    with pytest.raises(asyncio.CancelledError):
        await engine._operation(form["id"], "submit")
    assert store.job(form["id"])["state"] == "submission_uncertain"
    assert store.submissions_today() == 1


def test_excel_contains_paths_exact_answers_long_events_and_outcomes(store, vacancy, tmp_path):
    from openpyxl import load_workbook

    job = store.add(vacancy)
    value = "A" * 30000 + '=HYPERLINK("bad")'
    store.event(
        job["id"],
        "browser_action",
        {"field": "Motivation", "value": value, "evidence": "Verified " * 5000},
    )
    path = tmp_path / "cv.pdf"
    path.write_bytes(b"test attachment")
    store.file(job["id"], "cv", path)
    store.update(job["id"], state="hired")
    store.answer("Salary?", "10000 PLN", job["id"])
    book = load_workbook(io.BytesIO(workbook(store)))
    assert book["Applications"]["I2"].value == "hired"
    assert book["Files"]["C2"].value == str(path.resolve())
    parts = [row[4].value for row in list(book["Answers"].rows)[1:]]
    assert parts[0] == "A" * 30000 and parts[1] == '\'=HYPERLINK("bad")'
    history = "".join(row[5].value for row in list(book["History"].rows)[1:])
    assert json.loads(history)["value"] == value
    assert book["Saved answers"]["A2"].value == job["id"]
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)


@pytest.mark.asyncio
async def test_cycle_runs_rank_prepare_submit_and_persists_history(store, vacancy, monkeypatch):
    engine = Engine(store)
    store.save_settings(Settings(context=PROFILE, auto_submit=True))
    engine.search.discover = AsyncMock(return_value=([vacancy], []))
    engine.search.hydrate = AsyncMock(return_value=vacancy)
    engine.ranker.evaluate = AsyncMock(
        return_value={"score": 95, "confidence": 0.95, "category": "priority"}
    )
    engine.documents.prepare = AsyncMock(
        return_value={"profile": PROFILE, "revision": "test", "files": {}, "checks": {}}
    )

    async def apply(job, settings, store, allow_submit):
        assert allow_submit
        store.update(job["id"], state="submitted")
        store.event(job["id"], "submission_confirmed", {"text": "Thank you for applying"})

    engine.browser.apply = apply
    await engine._cycle()
    job = store.jobs()[0]
    assert job["state"] == "submitted"
    assert {event["stage"] for event in store.events()} >= {
        "cycle_started",
        "ranked",
        "prepared",
        "submission_confirmed",
        "cycle_finished",
    }
    await engine._cycle()
    engine.ranker.evaluate.assert_awaited_once()


@pytest.mark.asyncio
async def test_tracker_materializes_and_updates_workbook(store, vacancy, tmp_path):
    from openpyxl import load_workbook

    tracker = ExcelAgent(store)
    tracker.path = tmp_path / "applications.xlsx"
    job = store.add(vacancy)
    await tracker.flush()
    assert tracker.path.is_file() and not tracker.error
    store.update(job["id"], state="rejected")
    store.event(job["id"], "user_outcome", {"status": "rejected"})
    await tracker.flush()
    assert load_workbook(tracker.path)["Applications"]["I2"].value == "rejected"
    assert not Path(str(tracker.path) + ".tmp").exists()


@pytest.mark.asyncio
async def test_verified_pdf_upload_and_exact_answer_audit(store, form, tmp_path, monkeypatch):
    path = tmp_path / "cv.pdf"
    path.write_bytes(b"%PDF test attachment")
    store.file(form["id"], "cv", path)
    store.update(form["id"], package={"profile": PROFILE, "cv_path": str(path), "revision": "test"})
    plans = [
        {
            "tool": "type",
            "args": {"ref": "ref_1", "text": "B2+"},
            "field": "English level",
            "value": "B2+",
            "evidence": "B2+",
        },
        {"tool": "upload_file", "args": {"ref": "ref_1", "filePath": "/wrong.pdf"}},
        {"done": True},
    ]
    monkeypatch.setattr(browser, "call_ai", AsyncMock(side_effect=[json.dumps(p) for p in plans]))
    monkeypatch.setattr(browser.JevRanker, "supports_answer", AsyncMock(return_value=True))
    await browser.BrowserAgent().apply(store.job(form["id"]), store.settings(), store, False)
    calls = FakeMCP.instances[0].calls
    upload = next(args for name, args in calls if name == "upload_file")
    assert "filePath" not in upload and upload["file"].endswith(".pdf")
    assert any(name == "create_file" for name, _ in calls)
    recorded = [
        e
        for e in store.events(form["id"])
        if e["stage"] == "browser_action" and e["payload"].get("field")
    ]
    assert recorded[0]["payload"]["value"] == "B2+"
    assert store.job(form["id"])["state"] == "form_ready"


@pytest.mark.asyncio
async def test_modified_pdf_is_not_uploaded(store, form, tmp_path, monkeypatch):
    path = tmp_path / "cv.pdf"
    path.write_bytes(b"original")
    store.file(form["id"], "cv", path)
    path.write_bytes(b"changed")
    store.update(form["id"], package={"profile": PROFILE, "cv_path": str(path), "revision": "test"})
    monkeypatch.setattr(
        browser, "call_ai", AsyncMock(return_value='{"tool":"upload_file","args":{"ref":"ref_1"}}')
    )
    await browser.BrowserAgent().apply(store.job(form["id"]), store.settings(), store, False)
    assert store.job(form["id"])["state"] == "needs_review"
    assert not any(name == "upload_file" for name, _ in FakeMCP.instances[0].calls)


def test_api_import_export_outcome_and_missing_key(store, vacancy, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from careertwin.api.routes import automation

    store.save_settings(Settings(context=PROFILE, decision_backend="typesafe"))
    monkeypatch.setattr(automation, "engine", Engine(store))
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    app = FastAPI()
    app.include_router(automation.router)
    with TestClient(app) as client:
        base = "/v1/automation"
        assert len(client.get(base + "/sources").json()) == 30
        job = client.post(base + "/jobs", json=vacancy.model_dump()).json()
        assert client.post(base + "/run").status_code == 409
        assert (
            client.post(base + f"/jobs/{job['id']}/outcome", json={"status": "hired"}).json()[
                "state"
            ]
            == "hired"
        )
        assert (
            client.post(
                base + f"/jobs/{job['id']}/actions", json={"operation": "submit"}
            ).status_code
            == 409
        )
        assert client.get(base + "/export.xlsx").content.startswith(b"PK")


def remote_response():
    answers = {
        name: {"choice": "4", "probabilities": {str(i): float(i == 4) for i in range(5)}}
        for name in DIMENSIONS
    }
    answers.update(
        {
            name: {
                "choice": "satisfied",
                "probabilities": {
                    option: float(option == "satisfied")
                    for option in ("satisfied", "gap", "unknown", "not_required")
                },
            }
            for name in GATES
        }
    )
    return {"model": "cloud-test-model", "answers": answers}


@pytest.mark.asyncio
async def test_cloud_decisions_need_no_typesafe_key(vacancy, monkeypatch):
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    monkeypatch.setenv("SWITCHYARD_URL", "https://gateway.example.com")

    def handle(request):
        assert str(request.url) == "https://gateway.example.com/v1/decisions"
        body = json.loads(request.content)
        assert body["model"] == "cloud-selected-model"
        assert all(question["type"] == "choice" for question in body["questions"].values())
        assert set(body["questions"]["skills"]["criteria"]) == {str(i) for i in range(5)}
        return httpx.Response(200, json=remote_response())

    evaluator = JevRanker(httpx.MockTransport(handle))
    result = await evaluator.evaluate(
        vacancy, Settings(context=PROFILE, decision_model="cloud-selected-model")
    )
    assert result["category"] == "priority" and result["score"] == 100
    assert result["model"] == "cloud-test-model"
    assert result["backend"] == "switchyard"
    assert result["confidence_method"] == "model_estimate_uncalibrated"


@pytest.mark.asyncio
@pytest.mark.parametrize("bad", ["missing", "wrong_choice", "wrong_sum", "negative"])
async def test_invalid_cloud_distributions_fail_closed(vacancy, bad):
    payload = remote_response()
    answer = payload["answers"]["skills"]
    if bad == "missing":
        answer.pop("probabilities")
    elif bad == "wrong_choice":
        answer["choice"] = "1"
    elif bad == "wrong_sum":
        answer["probabilities"]["4"] = 0.2
    else:
        answer["probabilities"]["1"] = -0.1
    evaluator = JevRanker(httpx.MockTransport(lambda request: httpx.Response(200, json=payload)))
    with pytest.raises(ValueError, match="distribution"):
        await evaluator.evaluate(vacancy, Settings())


@pytest.mark.asyncio
async def test_cloud_uncertainty_and_expected_fit(vacancy):
    payload = remote_response()
    payload["answers"]["skills"] = {
        "choice": "4",
        "probabilities": {"0": 0, "1": 0, "2": 0, "3": 0.4, "4": 0.6},
    }
    evaluator = JevRanker(httpx.MockTransport(lambda request: httpx.Response(200, json=payload)))
    result = await evaluator.evaluate(vacancy, Settings())
    assert result["score"] == 96
    assert result["confidence"] == 0.6 and result["category"] == "needs_info"


def test_cloud_backend_is_default_without_jev_key(store, monkeypatch):
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    assert store.settings().decision_backend == "switchyard"
    assert Engine(store).status()["decision_configured"]


@pytest.mark.asyncio
async def test_cloud_service_error_preserves_actionable_reason(vacancy):
    evaluator = JevRanker(
        httpx.MockTransport(
            lambda request: httpx.Response(
                503, json={"error": {"message": "No decision model answered: provider rate limit"}}
            )
        )
    )
    with pytest.raises(RuntimeError, match="provider rate limit"):
        await evaluator.evaluate(vacancy, Settings())


@pytest.mark.asyncio
async def test_available_cloud_model_listing(monkeypatch):
    from careertwin.api.routes import automation

    class Client:
        def __init__(self, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            pass

        async def get(self, url, headers):
            assert url.endswith("/v1/models")
            return httpx.Response(
                200,
                request=httpx.Request("GET", url),
                json={"data": [{"id": "auto"}, {"id": "cloud/model"}, {"id": 123}]},
            )

    monkeypatch.setattr(automation.httpx, "AsyncClient", Client)
    assert await automation.decision_models() == [{"id": "auto"}, {"id": "cloud/model"}]


@pytest.mark.asyncio
async def test_salary_is_resolved_scoped_and_audited_before_typing(store, form, monkeypatch):
    question = "Expected monthly salary PLN gross"
    evidence = "12100 PLN / month (gross)"
    choice = {
        "question": question,
        "answer": evidence,
        "amount": 12100,
        "mode": "auto",
        "evidence": "10000–16000 PLN gross per month",
    }
    monkeypatch.setattr(browser, "expectation", AsyncMock(return_value=choice))
    monkeypatch.setattr(
        browser,
        "call_ai",
        AsyncMock(
            side_effect=[
                json.dumps({"questions": [question]}),
                json.dumps(
                    {
                        "tool": "type",
                        "args": {"ref": "ref_1", "text": "12100"},
                        "field": question,
                        "value": "12100",
                        "evidence": evidence,
                    }
                ),
                '{"done":true}',
            ]
        ),
    )
    monkeypatch.setattr(browser.JevRanker, "supports_answer", AsyncMock(return_value=True))
    await browser.BrowserAgent().apply(form, store.settings(), store, False)
    assert store.job(form["id"])["state"] == "form_ready"
    assert store.answers(form["id"])[question] == evidence
    assert store.answers() == {}
    events = store.events(form["id"])
    assert any(
        event["stage"] == "salary_expectation" and event["payload"] == choice for event in events
    )
    assert any(
        event["stage"] == "browser_action" and event["payload"].get("value") == "12100"
        for event in events
    )


@pytest.mark.asyncio
async def test_numeric_salary_without_budget_pauses_before_typing(store, form, monkeypatch):
    monkeypatch.setattr(browser, "expectation", AsyncMock(return_value=None))
    monkeypatch.setattr(
        browser,
        "call_ai",
        AsyncMock(
            return_value=json.dumps(
                {
                    "tool": "type",
                    "args": {"ref": "ref_1", "text": "120000"},
                    "field": "Expected annual salary",
                    "value": "120000",
                    "evidence": "some history",
                }
            )
        ),
    )
    await browser.BrowserAgent().apply(form, store.settings(), store, False)
    assert store.job(form["id"])["state"] == "awaiting_answers"
    assert not any(name == "type" for name, _ in FakeMCP.instances[0].calls)


# --- the form filler stays on the employer's site and never shows a model a secret ---------------

TYPE_B2 = {
    "tool": "type", "args": {"ref": "ref_1", "text": "B2+"},
    "field": "English level", "value": "B2+", "evidence": "B2+",
}
PASSWORD = "S3cret-passw0rd!"
LOGIN_PAGE = '[ref_1] textbox "Email"\n[ref_3] textbox "Password"\n[ref_2] button "Sign in"'
PASSWORD_PLAN = {
    "tool": "type", "args": {"ref": "ref_3", "text": "{{secret:ACCOUNT_PASSWORD}}"},
    "field": "Password", "value": "{{secret:ACCOUNT_PASSWORD}}",
}


def types(calls):
    return [args for name, args in calls if name == "type"]


@pytest.fixture
def secure(monkeypatch, tmp_path):
    """A login configured for example.com, and a kill switch that lives in the test's own folder."""
    from careertwin.automation import guard

    monkeypatch.setenv("ACCOUNT_EMAIL", "login@example.com")
    monkeypatch.setenv("ACCOUNT_PASSWORD", PASSWORD)
    monkeypatch.setenv("CAREERTWIN_LOGIN_DOMAINS", "example.com")
    monkeypatch.setattr(guard, "DATA_DIR", tmp_path / "guard")
    monkeypatch.setattr(FakeMCP, "page", LOGIN_PAGE)
    monkeypatch.setattr(browser.JevRanker, "supports_answer", AsyncMock(return_value=True))


@pytest.mark.asyncio
async def test_a_form_on_another_site_is_not_filled(store, form, monkeypatch):
    monkeypatch.setattr(FakeMCP, "tab_url", "https://attacker.example/apply")
    monkeypatch.setattr(browser, "call_ai", AsyncMock(return_value=json.dumps(TYPE_B2)))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=True)
    job = store.job(form["id"])
    assert job["state"] == "needs_review" and "outside the employer's site" in job["error"]
    assert store.events(form["id"])[-1]["stage"] == "domain_blocked"
    assert types(FakeMCP.instances[0].calls) == []


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "url",
    ["https://example.com/jobs/python", "https://careers.example.com/apply", "https://boards.greenhouse.io/example/jobs/1",
     "https://example.wd5.myworkdayjobs.com/en-US/job/1"],
)
async def test_the_employers_site_and_application_systems_are_filled(store, form, monkeypatch, url):
    monkeypatch.setattr(FakeMCP, "tab_url", url)
    monkeypatch.setattr(browser, "call_ai", AsyncMock(side_effect=[json.dumps(TYPE_B2), '{"done": true}']))
    monkeypatch.setattr(browser.JevRanker, "supports_answer", AsyncMock(return_value=True))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert types(FakeMCP.instances[0].calls) == [{"ref": "ref_1", "text": "B2+", "tabId": 123}]


@pytest.mark.asyncio
async def test_a_link_that_leads_off_site_stops_the_next_action(store, form, monkeypatch):
    """The page's own button sends the tab elsewhere; whatever the page then says, nothing is typed."""

    class Redirected(FakeMCP):
        async def tool(self, name, args):
            if name == "click":
                self.calls.append((name, args))
                self.__class__.tab_url = "https://attacker.example/steal"
                return {"content": [{"type": "text", "text": '[ref_1] textbox "Send us your CV and salary"'}]}
            return await super().tool(name, args)

    monkeypatch.setattr(browser, "MCP", Redirected)
    monkeypatch.setattr(Redirected, "tab_url", "https://example.com/jobs/python")
    monkeypatch.setattr(Redirected, "page", '[ref_1] textbox "English level"\n[ref_2] link "Next step"')
    plans = [{"tool": "click", "args": {"ref": "ref_2"}}, TYPE_B2]
    monkeypatch.setattr(browser, "call_ai", AsyncMock(side_effect=[json.dumps(p) for p in plans]))
    monkeypatch.setattr(browser.JevRanker, "supports_answer", AsyncMock(return_value=True))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert store.job(form["id"])["state"] == "needs_review"
    assert any(name == "click" for name, _ in Redirected.instances[0].calls)  # it did follow the link
    assert types(Redirected.instances[0].calls) == []


@pytest.mark.asyncio
async def test_the_agent_cannot_navigate_away_from_the_employers_site(store, form, monkeypatch):
    plan = {"tool": "navigate", "args": {"url": "https://attacker.example/collect"}}
    monkeypatch.setattr(browser, "call_ai", AsyncMock(return_value=json.dumps(plan)))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert store.job(form["id"])["state"] == "needs_review"
    assert not any(name == "navigate" for name, _ in FakeMCP.instances[0].calls)


@pytest.mark.asyncio
async def test_without_the_tabs_tool_the_agent_does_not_act(store, form, monkeypatch):
    class Old(FakeMCP):
        def __init__(self):
            super().__init__()
            del self.definitions["tabs_list"]

    monkeypatch.setattr(browser, "MCP", Old)
    monkeypatch.setattr(browser, "call_ai", AsyncMock(return_value=json.dumps(TYPE_B2)))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert "tabs_list" in store.job(form["id"])["error"] and types(Old.instances[0].calls) == []


@pytest.mark.asyncio
async def test_the_password_is_typed_by_the_executor_and_never_shown_to_a_model_or_stored(store, form, secure, monkeypatch):
    email = {"tool": "type", "args": {"ref": "ref_1", "text": "login@example.com"}, "field": "Email",
             "value": "login@example.com", "evidence": "login@example.com"}
    model = AsyncMock(side_effect=[json.dumps(email), json.dumps(PASSWORD_PLAN), '{"done": true}'])
    monkeypatch.setattr(browser, "call_ai", model)
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)

    typed = types(FakeMCP.instances[0].calls)
    assert [item["text"] for item in typed] == ["login@example.com", PASSWORD]  # the real value reached the page
    prompts = json.dumps([call.args for call in model.call_args_list])
    assert PASSWORD not in prompts and "{{secret:ACCOUNT_PASSWORD}}" in prompts  # the model only ever saw the placeholder
    assert PASSWORD not in json.dumps(store.events()) and "{{secret:ACCOUNT_PASSWORD}}" in json.dumps(store.events())
    assert store.job(form["id"])["state"] == "form_ready"


@pytest.mark.asyncio
async def test_the_login_is_offered_to_the_model_only_when_it_is_fully_configured(store, form, secure, monkeypatch):
    monkeypatch.setenv("CAREERTWIN_LOGIN_DOMAINS", "")
    model = AsyncMock(return_value='{"done": true}')
    monkeypatch.setattr(browser, "call_ai", model)
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert json.loads(model.call_args.args[1])["login"] is None


@pytest.mark.asyncio
async def test_a_page_that_echoes_the_password_is_masked_before_the_model_reads_it(store, form, secure, monkeypatch):
    class Echo(FakeMCP):
        async def tool(self, name, args):
            if name == "type":
                return {"content": [{"type": "text", "text": f'[ref_3] textbox "Password" value="{PASSWORD}"'}]}
            return await super().tool(name, args)

    monkeypatch.setattr(browser, "MCP", Echo)
    model = AsyncMock(side_effect=[json.dumps(PASSWORD_PLAN), '{"done": true}'])
    monkeypatch.setattr(browser, "call_ai", model)
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert PASSWORD not in json.dumps([call.args for call in model.call_args_list])
    assert PASSWORD not in json.dumps(store.events())


@pytest.mark.asyncio
async def test_the_password_is_not_typed_on_a_domain_that_is_not_a_login_domain(store, form, secure, monkeypatch):
    monkeypatch.setenv("CAREERTWIN_LOGIN_DOMAINS", "linkedin.com")
    monkeypatch.setattr(FakeMCP, "tab_url", "https://boards.greenhouse.io/example/jobs/1")  # a shared system
    monkeypatch.setattr(browser, "call_ai", AsyncMock(return_value=json.dumps(PASSWORD_PLAN)))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert store.job(form["id"])["state"] == "awaiting_answers"
    assert store.events(form["id"])[-1]["stage"] == "secret_blocked" and types(FakeMCP.instances[0].calls) == []


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "text",
    ["{{secret:CAREERTWIN_API_TOKEN}}", "my {{secret:ACCOUNT_PASSWORD}} here", "{{secret:ACCOUNT_PASSWORD}}{{secret:ACCOUNT_PASSWORD}}"],
)
async def test_only_the_account_password_and_only_on_its_own_may_be_typed(store, form, secure, monkeypatch, text):
    monkeypatch.setenv("CAREERTWIN_API_TOKEN", "t" * 40)
    plan = {**PASSWORD_PLAN, "args": {"ref": "ref_3", "text": text}, "value": text}
    monkeypatch.setattr(browser, "call_ai", AsyncMock(return_value=json.dumps(plan)))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert store.job(form["id"])["state"] == "needs_review" and types(FakeMCP.instances[0].calls) == []


@pytest.mark.asyncio
async def test_a_secret_is_not_typed_into_a_field_that_is_not_a_password_field(store, form, secure, monkeypatch):
    plan = {**PASSWORD_PLAN, "args": {"ref": "ref_1", "text": "{{secret:ACCOUNT_PASSWORD}}"}}  # ref_1 is the Email box
    monkeypatch.setattr(browser, "call_ai", AsyncMock(return_value=json.dumps(plan)))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert store.job(form["id"])["state"] == "needs_review" and types(FakeMCP.instances[0].calls) == []


@pytest.mark.asyncio
async def test_the_kill_switch_stops_the_filler_before_it_does_anything(store, form, secure, monkeypatch):
    from careertwin.automation import guard

    guard.stop("test")
    model = AsyncMock()
    monkeypatch.setattr(browser, "call_ai", model)
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=True)
    job = store.job(form["id"])
    assert job["state"] == "needs_review" and "kill switch" in job["error"]
    model.assert_not_called()
    assert FakeMCP.instances[0].calls == []


def test_events_never_keep_a_secret(store, monkeypatch):
    monkeypatch.setenv("ACCOUNT_PASSWORD", PASSWORD)
    store.event(None, "note", {"text": f"typed {PASSWORD} by mistake", "password": "hunter2", "nested": [{"k": f"a{PASSWORD}b"}]})
    payload = json.dumps(store.events()[-1]["payload"])
    assert PASSWORD not in payload and "hunter2" not in payload


def test_the_kill_switch_is_reachable_from_the_api(store, tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from careertwin.api.routes import automation
    from careertwin.automation import guard

    store.save_settings(Settings(context=PROFILE))
    monkeypatch.setattr(guard, "DATA_DIR", tmp_path / "guard")
    monkeypatch.setattr(automation, "engine", Engine(store))
    app = FastAPI()
    app.include_router(automation.router)
    with TestClient(app) as client:
        base = "/v1/automation"
        assert client.get(base + "/status").json()["stopped"] is False
        assert client.post(base + "/kill").json() == {"stopped": True}
        assert guard.stopped() and client.get(base + "/status").json()["stopped"] is True
        assert client.post(base + "/run").status_code == 423  # locked while the switch is on
        assert client.delete(base + "/kill").json() == {"stopped": False}
        assert not guard.stopped()


def test_a_search_cycle_does_not_start_while_the_kill_switch_is_on(store, tmp_path, monkeypatch):
    from careertwin.automation import guard

    store.save_settings(Settings(context=PROFILE))
    monkeypatch.setattr(guard, "DATA_DIR", tmp_path / "guard")
    guard.stop("test")
    with pytest.raises(guard.Stopped):
        Engine(store).run()


@pytest.mark.asyncio
async def test_a_one_time_code_is_read_from_the_mailbox_and_typed(store, form, monkeypatch):
    class Mailbox:
        asked = []

        def __init__(self, store):
            pass

        async def wait_for_code(self, hint):
            self.asked.append(hint)
            return "482913"

    monkeypatch.setattr(browser, "MailAgent", Mailbox)
    monkeypatch.setattr(browser, "mail_configured", lambda: True)
    monkeypatch.setattr(FakeMCP, "page", '[ref_1] textbox "Verification code"\n[ref_2] button "Continue"')
    code_plan = {"tool": "type", "args": {"ref": "ref_1", "text": "482913"}, "field": "Verification code", "value": "482913", "evidence": "482913"}
    model = AsyncMock(side_effect=[json.dumps({"need_code": "workday"}), json.dumps(code_plan), '{"done": true}'])
    monkeypatch.setattr(browser, "call_ai", model)
    monkeypatch.setattr(browser.JevRanker, "supports_answer", AsyncMock(return_value=True))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert Mailbox.asked == ["workday"] and types(FakeMCP.instances[0].calls)[0]["text"] == "482913"
    assert json.loads(model.call_args_list[1].args[1])["verification"] == {"code": "482913", "from": "workday"}
    assert store.job(form["id"])["state"] == "form_ready"


@pytest.mark.asyncio
async def test_without_a_code_the_filler_asks_the_user_instead_of_guessing(store, form, monkeypatch):
    monkeypatch.setattr(browser, "mail_configured", lambda: False)
    monkeypatch.setattr(browser, "call_ai", AsyncMock(return_value=json.dumps({"need_code": "workday"})))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    job = store.job(form["id"])
    assert job["state"] == "awaiting_answers" and "one-time code" in json.dumps(store.events(form["id"])[-1])
    assert types(FakeMCP.instances[0].calls) == []


@pytest.mark.asyncio
async def test_an_invented_code_is_not_typed(store, form, monkeypatch):
    monkeypatch.setattr(browser, "mail_configured", lambda: False)
    plan = {"tool": "type", "args": {"ref": "ref_1", "text": "123456"}, "field": "Code", "value": "123456", "evidence": "123456"}
    monkeypatch.setattr(browser, "call_ai", AsyncMock(return_value=json.dumps(plan)))
    await browser.BrowserAgent().apply(form, store.settings(), store, allow_submit=False)
    assert store.job(form["id"])["state"] == "needs_review" and types(FakeMCP.instances[0].calls) == []
