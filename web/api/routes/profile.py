"""CareerTwin API — profile routes."""

from __future__ import annotations

import asyncio
import json
import logging
import re
import sqlite3
from typing import Any
from uuid import uuid4

from fastapi import APIRouter, File, Request, UploadFile

from database import (
    UserProfileCreate,
    UserProfileUpdate,
    _db_conn,
    _row_to_dict,
)

import os

logger = logging.getLogger("careertwin.api")

router = APIRouter()


@router.get("/v1/profile/{username}")
async def get_profile(username: str) -> dict[str, Any]:
    with _db_conn() as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            "SELECT * FROM user_profiles WHERE username = ?", (username,)
        ).fetchone()
    if row is None:
        return {"exists": False, "username": username}
    return {"exists": True, **_row_to_dict(row)}


@router.delete("/v1/profile/{username}")
async def delete_profile(username: str) -> dict[str, Any]:
    with _db_conn() as conn:
        conn.execute("DELETE FROM user_profiles WHERE username = ?", (username,))
        conn.commit()
    return {"success": True}


@router.post("/v1/profile")
async def create_or_update_profile(profile: UserProfileCreate) -> dict[str, Any]:
    with _db_conn() as conn:
        existing = conn.execute(
            "SELECT id FROM user_profiles WHERE username = ?", (profile.username,)
        ).fetchone()

        if existing:
            conn.execute("""
                UPDATE user_profiles SET
                    full_name = ?, email = ?, phone = ?, location = ?,
                    github = ?, linkedin = ?, website = ?, summary = ?,
                    skills = ?, updated_at = CURRENT_TIMESTAMP
                WHERE username = ?
            """, (
                profile.full_name, profile.email, profile.phone, profile.location,
                profile.github, profile.linkedin, profile.website, profile.summary,
                json.dumps(profile.skills), profile.username,
            ))
        else:
            conn.execute("""
                INSERT INTO user_profiles (id, username, full_name, email, phone, location, github, linkedin, website, summary, skills)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                uuid4().hex[:12], profile.username, profile.full_name, profile.email,
                profile.phone, profile.location, profile.github, profile.linkedin,
                profile.website, profile.summary, json.dumps(profile.skills),
            ))

        conn.commit()
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            "SELECT * FROM user_profiles WHERE username = ?", (profile.username,)
        ).fetchone()
    return {"success": True, **_row_to_dict(row)}


@router.patch("/v1/profile/{username}")
async def update_profile_fields(username: str, update: UserProfileUpdate) -> dict[str, Any]:
    with _db_conn() as conn:
        existing = conn.execute(
            "SELECT id FROM user_profiles WHERE username = ?", (username,)
        ).fetchone()
        if not existing:
            return {"success": False, "error": "Profile not found"}

        fields: list[str] = []
        values: list[Any] = []
        for field_name in ["full_name", "email", "phone", "location", "github", "linkedin", "website", "summary"]:
            val = getattr(update, field_name)
            if val is not None:
                fields.append(f"{field_name} = ?")
                values.append(val)
        if update.skills is not None:
            fields.append("skills = ?")
            values.append(json.dumps(update.skills))

        if fields:
            fields.append("updated_at = CURRENT_TIMESTAMP")
            values.append(username)
            conn.execute(
                f"UPDATE user_profiles SET {', '.join(fields)} WHERE username = ?",
                values,
            )
            conn.commit()

        conn.row_factory = sqlite3.Row
        row = conn.execute(
            "SELECT * FROM user_profiles WHERE username = ?", (username,)
        ).fetchone()
    return {"success": True, **_row_to_dict(row)}



@router.post("/v1/profile/upload-resume")
async def upload_resume(
    request: Request,
    file: UploadFile = File(...),
    username: str = "default",
) -> dict[str, Any]:
    """Upload a PDF/DOCX/TXT resume, extract profile data, and return it for scoring.

    If the client disconnects (e.g. the user picks a different file mid-upload
    and aborts the request), cancel the in-flight task so we stop waiting on
    the LLM provider instead of burning tokens/compute on a discarded result.
    """
    task = asyncio.ensure_future(_process_resume_upload(file, username))

    async def _watch_disconnect() -> None:
        while not task.done():
            if await request.is_disconnected():
                task.cancel()
                return
            await asyncio.sleep(0.5)

    watcher = asyncio.ensure_future(_watch_disconnect())
    try:
        return await task
    except asyncio.CancelledError:
        return {"success": False, "error": "Upload cancelled"}
    finally:
        watcher.cancel()


async def _process_resume_upload(file, username: str) -> dict[str, Any]:
    """Internal handler for resume upload processing.
    
    Uses LLM-based extraction for any CV format, with regex fallback.
    """
    import tempfile
    from pathlib import Path

    if file is None:
        return {"success": False, "error": "No file uploaded"}

    filename = getattr(file, "filename", "") or ""
    suffix = Path(filename).suffix or ".txt"

    try:
        content = await file.read()
        if not content:
            return {"success": False, "error": "Empty file"}

        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp.write(content)
            tmp_path = Path(tmp.name)

        try:
            raw_text = _extract_text_from_file(tmp_path, suffix)

            llm_result, llm_info = await _llm_extract_profile(raw_text, filename)

            if llm_result and llm_result.get("fullName"):
                profile_data = llm_result
                profile_data["_provider"] = llm_info
            else:
                logger.info("LLM extraction unavailable (%s), using regex fallback", llm_info)
                from careertwin.collectors.resume import ResumeParser
                parser = ResumeParser()
                graph = parser.parse(tmp_path)
                profile_data = _graph_to_profile(graph, raw_text, tmp_path, suffix)
                profile_data["_provider"] = "regex-fallback"

            if username and profile_data.get("fullName"):
                _update_profile_from_resume(username, profile_data)

            provider = profile_data.pop("_provider", "unknown")
            msg = _build_resume_message(profile_data)
            if provider == "regex-fallback":
                msg += " (regex fallback — set GROQ_API_KEY for AI-powered extraction)"

            return {
                "success": True,
                "profile": profile_data,
                "message": msg,
                "provider": provider,
            }
        finally:
            tmp_path.unlink(missing_ok=True)
    except Exception as e:
        logger.error("Resume parsing failed: %s", e)
        return {"success": False, "error": f"Failed to parse resume: {e}"}


def _extract_text_from_file(path: Path, suffix: str) -> str:
    """Extract raw text from PDF, DOCX, or TXT."""
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".pdf":
        texts = []
        try:
            from pdfminer.high_level import extract_text
            t = extract_text(str(path))
            if t:
                texts.append(t)
        except Exception:
            pass
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            raw = ""
            for page in reader.pages:
                raw += (page.extract_text() or "") + "\n"
            if raw.strip():
                texts.append(raw)
        except Exception:
            pass
        return max(texts, key=len) if texts else ""

    if suffix in (".docx", ".doc"):
        try:
            from docx import Document
            doc = Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        text += "\n" + " | ".join(cells)
            return text
        except Exception:
            pass

    return ""


async def _llm_extract_profile(text: str, filename: str) -> tuple[dict[str, Any] | None, str]:
    """Use LLM to extract structured profile from ANY CV format.
    
    Returns (profile_dict, provider_used) or (None, error_message).
    """
    if not text or len(text) < 50:
        return None, "Resume text too short to parse"

    from ai import call_ai, parse_json_object, PROVIDERS

    text_chunk = text[:8000]

    system_prompt = """You are an expert resume/CV parser. Extract ALL structured data from the resume text.

Return a JSON object:
{
  "fullName": "Full Name",
  "title": "Current or most recent professional title, e.g. Senior Backend Engineer",
  "email": "email or empty",
  "phone": "phone or empty",
  "location": "City, Country or empty",
  "linkedin": "LinkedIn profile URL or empty",
  "github": "GitHub profile URL or empty",
  "summary": "professional summary or empty",
  "skills": ["skill1", "skill2"],
  "languages": ["English (Native)", "Russian (Fluent)"],
  "experience": [{"position":"Title","company":"Company","description":"What they did","start":"YYYY","end":"YYYY or Present"}],
  "education": [{"school":"University","degree":"Degree","field":"Field","start":"YYYY","end":"YYYY or Present"}],
  "certificates": [{"name":"Course/Cert Name","issuer":"Organization"}],
  "projects": [{"name":"Project","description":"Brief desc"}]
}

CRITICAL RULES:
- Title: use the candidate's most recent job title if stated, exactly as one clean title (e.g. "Backend Engineer"). If no title is stated anywhere, infer ONE concise, realistic, natural-sounding title (e.g. "Senior Backend Engineer") from their most recent experience and skills. NEVER combine multiple roles or specialties into one title with a slash or "and" (e.g. NOT "Backend Developer / AI Engineer", NOT "Python Developer and Data Analyst") — pick the single best-fitting title only. Never leave it empty unless there is truly nothing to infer from.
- Experience: each job gets its OWN entry. Date ranges go in start/end, NOT in position.
- Education: ONLY university degrees. Courses/bootcamps → certificates. Include the start/end year if the resume states them.
- LinkedIn: extract the full profile URL exactly as written (e.g. linkedin.com/in/username), or empty if not present.
- GitHub: extract the full profile URL exactly as written (e.g. github.com/username), or empty if not present.
- Skills: ONLY concrete, nameable technologies, tools, languages, frameworks, or platforms (e.g. "Python", "Kubernetes", "PostgreSQL", "React"). Do NOT include vague activities, processes, or responsibilities as skills (e.g. "deployment", "async processing", "collaboration", "problem solving", "communication") — those belong in the job description text, not the skills list.
- Languages: spoken/written languages with proficiency if stated, e.g. "English (Native)" — NOT programming languages (those go in skills).
- Return ONLY valid JSON, no markdown, no explanation."""

    providers_to_try = ["lmstudio", "groq", "openai", "openrouter"]
    last_error = ""

    for prov in providers_to_try:
        try:
            cfg = PROVIDERS.get(prov, {})
            key_env = cfg.get("key_env", "")
            if key_env and not os.getenv(key_env, ""):
                continue

            raw = await call_ai(
                system_prompt,
                f"Resume: {filename}\n\n{text_chunk}",
                provider=prov,
                model="",
                max_tokens=8000,
                temperature=0.1,
            )
            if not raw or len(raw) < 50:
                last_error = f"{prov}: empty response"
                continue

            result = parse_json_object(raw)
            if not result or not result.get("fullName"):
                last_error = f"{prov}: invalid JSON or missing name"
                continue

            profile = {
                "fullName": str(result.get("fullName", "")).strip(),
                "title": str(result.get("title", "")).strip(),
                "email": str(result.get("email", "")).strip(),
                "phone": str(result.get("phone", "")).strip(),
                "location": str(result.get("location", "")).strip(),
                "linkedin": str(result.get("linkedin", "")).strip(),
                "github": str(result.get("github", "")).strip(),
                "summary": str(result.get("summary", "")).strip(),
                "skills": _normalize_skills(result.get("skills", [])),
                "languages": _normalize_skills(result.get("languages", [])),
                "experience": _normalize_experience(result.get("experience", [])),
                "education": _normalize_education(result.get("education", [])),
                "certificates": _normalize_certificates(result.get("certificates", [])),
                "projects": _normalize_projects(result.get("projects", [])),
            }

            return profile, prov

        except Exception as e:
            last_error = f"{prov}: {e}"
            continue

    return None, f"No AI provider available. Set GROQ_API_KEY or OPENAI_API_KEY in .env. Last error: {last_error}"


def _normalize_skills(skills: Any) -> list[str]:
    if not isinstance(skills, list):
        return []
    return [str(s).strip() for s in skills if s and str(s).strip()][:50]


def _normalize_experience(exp: Any) -> list[dict]:
    if not isinstance(exp, list):
        return []
    result = []
    for item in exp:
        if not isinstance(item, dict):
            continue
        position = str(item.get("position", "")).strip()
        company = str(item.get("company", "")).strip()
        if not position and not company:
            continue
        desc = str(item.get("description", "")).strip()
        desc_lines = [l.strip() for l in desc.split("\n") if l.strip()]
        desc = "\n".join(desc_lines[:6])
        result.append({
            "position": position[:150],
            "company": company[:150],
            "description": desc[:600],
            "start": str(item.get("start", "")).strip(),
            "end": str(item.get("end", "")).strip(),
        })
    return result[:20]


def _normalize_education(edu: Any) -> list[dict]:
    if not isinstance(edu, list):
        return []
    result = []
    for item in edu:
        if not isinstance(item, dict):
            continue
        school = str(item.get("school", "")).strip()
        if not school:
            continue
        result.append({
            "school": school[:200],
            "degree": str(item.get("degree", "")).strip()[:100],
            "field": str(item.get("field", "")).strip()[:100],
            "start": str(item.get("start", "")).strip()[:20],
            "end": str(item.get("end", "")).strip()[:20],
        })
    return result[:10]


def _normalize_certificates(certs: Any) -> list[dict]:
    if not isinstance(certs, list):
        return []
    result = []
    for item in certs:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name", "")).strip()
        if not name:
            continue
        result.append({
            "name": name[:200],
            "issuer": str(item.get("issuer", "")).strip()[:100],
        })
    return result[:20]


def _normalize_projects(projects: Any) -> list[dict]:
    if not isinstance(projects, list):
        return []
    result = []
    for item in projects:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name", "")).strip()
        if not name:
            continue
        result.append({
            "name": name[:200],
            "description": str(item.get("description", "")).strip()[:400],
        })
    return result[:15]


def _graph_to_profile(graph, raw_text: str, tmp_path: Path, suffix: str) -> dict:
    """Convert CareerGraph from regex parser to profile dict (fallback)."""
    from careertwin.collectors.resume import ResumeParser

    skills = [s.name for s in graph.skills]
    experience = []
    for exp in graph.experience:
        position = (exp.position or "").strip()
        company = (exp.company or "").strip()
        desc = exp.description or ""

        if not position or len(position) < 3:
            continue
        import re as _re
        if _re.match(r"^\d{4}\s*[-–—]\s*(?:\d{4}|present)$", position.lower()):
            desc_lines = [l.strip() for l in desc.split("\n") if l.strip()]
            if desc_lines:
                candidate = desc_lines[0]
                if candidate and len(candidate) < 80 and not _re.match(r"^\d{4}", candidate):
                    position = candidate
                    desc = "\n".join(desc_lines[1:])
            if _re.match(r"^\d{4}\s*[-–—]\s*(?:\d{4}|present)$", position.lower()):
                continue
        if position.lower().startswith(("built ", "designed ", "developed ", "implemented ")):
            continue

        desc_lines = [l.strip() for l in desc.split("\n") if l.strip()]
        experience.append({
            "position": position[:120],
            "company": company[:120],
            "description": "\n".join(desc_lines[:8])[:500],
            "start": "",
            "end": "",
        })

    education = []
    certificates = []
    cert_kw = ["course", "certificate", "bootcamp", "workshop", "udemy", "coursera", "epam"]
    uni_kw = ["university", "institute", "college", "school", "academy", "bachelor", "master"]

    for edu in graph.education:
        inst = (edu.institution or "").lower()
        combined = f"{inst} {(edu.degree or '').lower()} {(edu.field or '').lower()}"
        is_cert = any(k in combined for k in cert_kw)
        is_uni = any(k in combined for k in uni_kw)
        if is_cert and not is_uni:
            certificates.append({
                "name": edu.degree or edu.institution or "",
                "issuer": edu.institution if edu.degree else "",
            })
        else:
            education.append({
                "school": edu.institution,
                "degree": edu.degree,
                "field": edu.field or "",
            })

    parser = ResumeParser()
    try:
        lines = [l.strip() for l in raw_text.split("\n") if l.strip()]
        sections = parser._split_sections(lines)
        raw_certs = parser._extract_certificates(sections)
        for c in raw_certs:
            name = c.get("name", "")
            if name and not any(x["name"] == name for x in certificates):
                certificates.append({"name": name, "issuer": c.get("issuer", "")})
    except Exception:
        pass

    if len(skills) < 3:
        text_skills = parser._extract_skills_from_text(raw_text)
        for s in text_skills:
            if s.name not in skills:
                skills.append(s.name)

    return {
        "fullName": graph.profile.full_name or "",
        "email": graph.profile.email or "",
        "phone": graph.profile.phone or "",
        "location": graph.profile.location or "",
        "summary": graph.profile.summary or "",
        "skills": skills[:50],
        "experience": experience[:20],
        "education": education[:10],
        "certificates": certificates[:20],
        "projects": [],
    }


def _update_profile_from_resume(username: str, profile_data: dict) -> None:
    """Update user profile in DB with extracted resume data."""
    try:
        with _db_conn() as conn:
            existing = conn.execute(
                "SELECT id FROM user_profiles WHERE username = ?", (username,)
            ).fetchone()
            if existing:
                conn.execute("""
                    UPDATE user_profiles SET
                        full_name = COALESCE(NULLIF(?, ''), full_name),
                        email = COALESCE(NULLIF(?, ''), email),
                        phone = COALESCE(NULLIF(?, ''), phone),
                        location = COALESCE(NULLIF(?, ''), location),
                        summary = COALESCE(NULLIF(?, ''), summary),
                        skills = COALESCE(NULLIF(?, '[]'), skills),
                        updated_at = CURRENT_TIMESTAMP
                    WHERE username = ?
                """, (
                    profile_data.get("fullName", ""),
                    profile_data.get("email", ""),
                    profile_data.get("phone", ""),
                    profile_data.get("location", ""),
                    profile_data.get("summary", ""),
                    json.dumps(profile_data.get("skills", [])),
                    username,
                ))
                conn.commit()
    except Exception as e:
        logger.warning("Failed to update profile from resume: %s", e)


def _build_resume_message(profile_data: dict) -> str:
    """Build human-readable summary of extracted data."""
    parts = []
    skills = profile_data.get("skills", [])
    exp = profile_data.get("experience", [])
    edu = profile_data.get("education", [])
    certs = profile_data.get("certificates", [])

    if skills:
        parts.append(f"{len(skills)} skills")
    if exp:
        parts.append(f"{len(exp)} experiences")
    if edu:
        parts.append(f"{len(edu)} education entries")
    if certs:
        parts.append(f"{len(certs)} certificates")

    return f"Resume parsed: {', '.join(parts)}" if parts else "Resume parsed successfully"
